from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

from docx import Document
from sqlalchemy.orm import Session

from ..models import Plan
from .sync import sync_single_plan


class DocEditError(RuntimeError):
    pass


def update_plan_cell(session: Session, plan_id: int, table_index: int, row_index: int, cell_index: int, text: str):
    plan = session.query(Plan).filter(Plan.id == plan_id).one()
    doc_path = Path(plan.source_doc_path).resolve()
    if not doc_path.exists():
        raise DocEditError(f"源文件不存在，无法回写: {doc_path.name}")

    temp_dir = Path(tempfile.mkdtemp(prefix="plan_edit_"))
    temp_doc_path = temp_dir / "editing.doc"
    temp_docx_path = temp_dir / "editing.docx"
    temp_output_doc = temp_dir / "edited_output.doc"
    shutil.copy2(doc_path, temp_doc_path)

    try:
        import pythoncom
        import win32com.client
    except Exception as exc:  # pragma: no cover
        raise DocEditError(f"缺少 pywin32 或 Word COM 环境: {exc}") from exc

    try:
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(str(temp_doc_path), False, True)
        document.SaveAs2(str(temp_docx_path), FileFormat=16)
    except Exception as exc:  # pragma: no cover
        raise DocEditError(f"将 .doc 转为临时 .docx 失败: {exc}") from exc
    finally:
        if "document" in locals() and document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if "word" in locals() and word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass

    try:
        docx_document = Document(str(temp_docx_path))
        docx_document.tables[table_index].cell(row_index, cell_index).text = str(text)
        docx_document.save(str(temp_docx_path))
    except Exception as exc:  # pragma: no cover
        raise DocEditError(f"修改临时 .docx 单元格失败: {exc}") from exc

    try:
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(str(temp_docx_path), False, False)
        document.SaveAs2(str(temp_output_doc), FileFormat=0)
    except Exception as exc:  # pragma: no cover
        raise DocEditError(f"将编辑后的 .docx 回写为 .doc 失败: {exc}") from exc
    finally:
        if "document" in locals() and document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if "word" in locals() and word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass

    shutil.copy2(temp_output_doc, doc_path)
    shutil.rmtree(temp_dir, ignore_errors=True)

    return sync_single_plan(session, doc_path)
