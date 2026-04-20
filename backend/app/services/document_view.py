from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from sqlalchemy.orm import Session, selectinload

from ..models import Plan, PlanBlock, PlanBlockCell, PlanBlockImage, PlanImage, PlanPage
from ..schemas import DocumentBlock, DocumentNavItem, DocumentPage, DocumentTableCell, PlanDocumentResponse


NAVIGATION = [
    {
        "key": "basic",
        "title": "基本情况",
        "children": [
            {"key": "basic_unit", "title": "单位基本情况"},
            {"key": "basic_facilities", "title": "单位主要消防设施"},
            {"key": "basic_key_parts", "title": "重点部位情况"},
            {"key": "basic_tips", "title": "重点提示"},
        ],
    },
    {
        "key": "images",
        "title": "图像资料",
        "children": [
            {"key": "images_route", "title": "行车路线图"},
            {"key": "images_site", "title": "单位总平面图"},
            {"key": "images_water", "title": "水源图"},
            {"key": "images_floor", "title": "楼层平面图"},
        ],
    },
    {
        "key": "assist",
        "title": "辅助决策",
        "children": [
            {"key": "assist_scenario", "title": "灾情假设"},
            {"key": "assist_communication", "title": "现场通信"},
            {"key": "assist_initial", "title": "首战力量部署图"},
            {"key": "assist_reinforce", "title": "增援力量部署图"},
        ],
    },
]


def _normalize(value: str) -> str:
    return " ".join(value.replace("\x07", " ").replace("\n", " ").split()).strip()


def _table_cell_text(tc) -> str:
    texts = [node.text or "" for node in tc.xpath(".//w:t")]
    return _normalize("".join(texts))


def _is_editable_cell(text: str, colspan: int, rowspan: int) -> int:
    blocked = {
        "基本情况",
        "单位基本情况",
        "单位内部主要消防设施",
        "重点部位情况",
        "重点提示",
        "风险评估",
        "灾情假设",
        "现场通信",
        "首战力量部署",
        "增援力量部署",
    }
    if not text.strip():
        return 0
    if rowspan > 1:
        return 0
    if colspan > 4:
        return 0
    if text.strip() in blocked:
        return 0
    return 1


def _apply_value_cell_editability(row_cells: list[DocumentTableCell]) -> None:
    meaningful = [(idx, cell) for idx, cell in enumerate(row_cells) if cell.text.strip()]
    if not meaningful:
        return

    # Single long text cells are usually free-form content and should remain editable.
    if len(meaningful) == 1:
        idx, cell = meaningful[0]
        cell.is_editable = 1 if len(cell.text.strip()) > 8 else 0
        return

    start = 0
    first_cell = meaningful[0][1]
    if first_cell.rowspan > 1:
        first_cell.is_editable = 0
        start = 1

    for pos, (idx, cell) in enumerate(meaningful[start:], start=0):
        # Alternate label/value pairs: value cells are every second logical item.
        cell.is_editable = 1 if pos % 2 == 1 else 0

    # Keep empty placeholders and structural cells readonly.
    for cell in row_cells:
        if not cell.text.strip():
            cell.is_editable = 0


def _display_table(table, table_index: int) -> list[list[DocumentTableCell]]:
    rendered: list[list[DocumentTableCell]] = []
    active_merges: dict[int, DocumentTableCell] = {}

    for row_index, tr in enumerate(table._tbl.tr_lst):
        row_cells: list[DocumentTableCell] = []
        col_index = 0
        cell_order = 0
        for tc in tr.tc_lst:
            while col_index in active_merges and tc.tcPr is not None and tc.tcPr.vMerge is None:
                active_merges.pop(col_index, None)
                col_index += 1

            text = _table_cell_text(tc)
            colspan = int(tc.tcPr.gridSpan.val) if tc.tcPr is not None and tc.tcPr.gridSpan is not None else 1
            vmerge = tc.tcPr.vMerge.val if tc.tcPr is not None and tc.tcPr.vMerge is not None else None

            if vmerge == "continue":
                merged_cell = active_merges.get(col_index)
                if merged_cell is not None:
                    merged_cell.rowspan += 1
                col_index += colspan
                continue

            cell = DocumentTableCell(
                text=text,
                colspan=colspan,
                rowspan=1,
                block_id=None,
                cell_order=cell_order,
                table_index=table_index,
                row_index=row_index,
                cell_index=col_index,
                is_editable=_is_editable_cell(text, colspan, 1),
            )
            row_cells.append(cell)

            if tc.tcPr is not None and tc.tcPr.vMerge is not None:
                for offset in range(colspan):
                    active_merges[col_index + offset] = cell
            else:
                for offset in range(colspan):
                    active_merges.pop(col_index + offset, None)

            col_index += colspan
            cell_order += 1

        if row_cells:
            _apply_value_cell_editability(row_cells)
            rendered.append(row_cells)

    return rendered


def _rows_to_block(title: str, rows: list[list[DocumentTableCell]]) -> DocumentBlock:
    return DocumentBlock(type="table", title=title, rows=rows)


def _table_slice(rows: list[list[str]], start_keywords: tuple[str, ...], end_keywords: tuple[str, ...]) -> list[list[str]]:
    start_index = None
    end_index = len(rows)
    for idx, row in enumerate(rows):
        first = _normalize(row[0]) if row else ""
        if start_index is None and any(keyword in first for keyword in start_keywords):
            start_index = idx
            continue
        if start_index is not None and any(keyword in first for keyword in end_keywords):
            end_index = idx
            break
    if start_index is None:
        return []
    return rows[start_index:end_index]


def _table_slice_indices(rows: list[list[str]], start_keywords: tuple[str, ...], end_keywords: tuple[str, ...]) -> tuple[int | None, int]:
    start_index = None
    end_index = len(rows)
    for idx, row in enumerate(rows):
        joined = " ".join(_normalize(cell) for cell in row if cell)
        if start_index is None and any(keyword in joined for keyword in start_keywords):
            start_index = idx
            continue
        if start_index is not None and end_keywords and any(keyword in joined for keyword in end_keywords):
            end_index = idx
            break
    return start_index, end_index


def _image_block(title: str, images: list[PlanImage] | list[PlanBlockImage]) -> DocumentBlock:
    return DocumentBlock(type="image_gallery", title=title, images=list(images))


def _rows_to_plain_text(rows: list[list[DocumentTableCell]]) -> list[list[str]]:
    return [[cell.text for cell in row] for row in rows]


def _match_image_category(title: str) -> str | None:
    normalized = _normalize(title)
    if not normalized:
        return None
    if "路线图" in normalized:
        return "行车路线图"
    if "总平面图" in normalized:
        return "单位总平面图"
    if "水源" in normalized:
        return "水源图"
    if "平面图" in normalized:
        return "楼层平面图"
    if "初战部署图" in normalized or "首战力量部署图" in normalized:
        return "首战力量部署图"
    if "增援力量部署图" in normalized:
        return "增援力量部署图"
    return None


def _extract_image_categories(docx_path: Path, images: list[PlanImage]) -> dict[str, list[PlanImage]]:
    categories = {
        "行车路线图": [],
        "单位总平面图": [],
        "水源图": [],
        "楼层平面图": [],
        "首战力量部署图": [],
        "增援力量部署图": [],
    }
    image_by_media = {image.doc_media_path: image for image in images}

    with ZipFile(docx_path) as archive:
        document_root = ET.fromstring(archive.read("word/document.xml"))
        rel_root = ET.fromstring(archive.read("word/_rels/document.xml.rels"))

        rel_map = {}
        for rel in rel_root:
            rel_id = rel.attrib.get("Id")
            target = rel.attrib.get("Target", "")
            if rel_id and "image" in rel.attrib.get("Type", ""):
                rel_map[rel_id] = f"word/{target}"

        ns = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "v": "urn:schemas-microsoft-com:vml",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        }
        r_ns = "{%s}id" % ns["r"]

        pending_category: str | None = None

        def filtered_media_paths(rids: list[str]) -> list[str]:
            return [rel_map[rid] for rid in rids if rid in rel_map]

        body = document_root.find("w:body", ns)
        if body is None:
            return categories

        for child in list(body):
            tag = child.tag.split("}")[-1]
            if tag == "p":
                text = _normalize("".join(node.text or "" for node in child.findall(".//w:t", ns)))
                rids = [node.attrib.get(r_ns) for node in child.findall(".//v:imagedata", ns)]
                rids = [rid for rid in rids if rid]
                category = _match_image_category(text)
                if rids and pending_category:
                    for media_path in filtered_media_paths(rids):
                        image = image_by_media.get(media_path)
                        if image and image not in categories[pending_category]:
                            categories[pending_category].append(image)
                    pending_category = None
                    if category:
                        pending_category = category
                    continue
                if category:
                    pending_category = category
                continue

            if tag == "tbl":
                for tr in child.findall("./w:tr", ns):
                    row_text = _normalize(
                        " ".join(_normalize("".join(node.text or "" for node in tc.findall(".//w:t", ns))) for tc in tr.findall("./w:tc", ns))
                    )
                    category = _match_image_category(row_text)
                    rids = [node.attrib.get(r_ns) for node in tr.findall(".//v:imagedata", ns)]
                    rids = [rid for rid in rids if rid]
                    if category and rids:
                        for media_path in filtered_media_paths(rids):
                            image = image_by_media.get(media_path)
                            if image and image not in categories[category]:
                                categories[category].append(image)

    return categories


def build_plan_document_from_source(plan: Plan, images: list[PlanImage]) -> PlanDocumentResponse:
    if not plan.source_docx_path:
        raise FileNotFoundError("该预案未生成 docx 缓存")

    docx_path = Path(plan.source_docx_path)
    document = Document(str(docx_path))
    tables = [_display_table(table, table_index) for table_index, table in enumerate(document.tables)]
    basic_table = tables[0] if tables else []
    assist_table = tables[1] if len(tables) > 1 else []
    image_categories = _extract_image_categories(docx_path, images)

    basic_table_plain = _rows_to_plain_text(basic_table)
    assist_table_plain = _rows_to_plain_text(assist_table)

    basic_unit_start, basic_unit_end = _table_slice_indices(basic_table_plain, ("单位基本情况",), ("单位内部主要消防设施",))
    basic_facility_start, basic_facility_end = _table_slice_indices(basic_table_plain, ("单位内部主要消防设施",), ("重点部位情况",))
    basic_key_start, basic_key_end = _table_slice_indices(basic_table_plain, ("重点部位情况",), ("重点提示",))
    basic_tip_start, basic_tip_end = _table_slice_indices(basic_table_plain, ("重点提示",), ())

    basic_unit_rows = basic_table[basic_unit_start:basic_unit_end] if basic_unit_start is not None else []
    basic_facility_rows = basic_table[basic_facility_start:basic_facility_end] if basic_facility_start is not None else []
    basic_key_rows = basic_table[basic_key_start:basic_key_end] if basic_key_start is not None else []
    basic_tip_rows = basic_table[basic_tip_start:basic_tip_end] if basic_tip_start is not None else []

    scenario_plain = _table_slice(assist_table_plain, ("灾情假设",), ("现场通信",))
    communication_plain = _table_slice(assist_table_plain, ("现场通信",), ("初战力量编成", "首战力量部署图", "初战部署图"))
    initial_plain = _table_slice(assist_table_plain, ("初战力量编成", "初战部署图"), ("增援力量编成",))
    reinforce_plain = _table_slice(assist_table_plain, ("增援力量编成", "增援部署图"), ())
    risk_plain = assist_table_plain[:4] if len(assist_table_plain) >= 4 else assist_table_plain

    scenario_rows = [row for row in assist_table if [cell.text for cell in row] in scenario_plain]
    communication_rows = [row for row in assist_table if [cell.text for cell in row] in communication_plain]
    initial_rows = [row for row in assist_table if [cell.text for cell in row] in initial_plain]
    reinforce_rows = [row for row in assist_table if [cell.text for cell in row] in reinforce_plain]
    risk_rows = assist_table[:len(risk_plain)]

    pages = [
        DocumentPage(
            key="basic",
            title="基本情况",
            blocks=[
                _rows_to_block("单位基本情况", basic_unit_rows),
                _rows_to_block("单位主要消防设施", basic_facility_rows),
                _rows_to_block("重点部位情况", basic_key_rows),
                _rows_to_block("重点提示", basic_tip_rows),
            ],
        ),
        DocumentPage(key="basic_unit", title="单位基本情况", parent_key="basic", blocks=[_rows_to_block("单位基本情况", basic_unit_rows)]),
        DocumentPage(key="basic_facilities", title="单位主要消防设施", parent_key="basic", blocks=[_rows_to_block("单位主要消防设施", basic_facility_rows)]),
        DocumentPage(key="basic_key_parts", title="重点部位情况", parent_key="basic", blocks=[_rows_to_block("重点部位情况", basic_key_rows)]),
        DocumentPage(key="basic_tips", title="重点提示", parent_key="basic", blocks=[_rows_to_block("重点提示", basic_tip_rows)]),
        DocumentPage(
            key="images",
            title="图像资料",
            blocks=[
                _image_block("行车路线图", image_categories["行车路线图"]),
                _image_block("单位总平面图", image_categories["单位总平面图"]),
                _image_block("水源图", image_categories["水源图"]),
                _image_block("楼层平面图", image_categories["楼层平面图"]),
            ],
        ),
        DocumentPage(key="images_route", title="行车路线图", parent_key="images", blocks=[_image_block("行车路线图", image_categories["行车路线图"])]),
        DocumentPage(key="images_site", title="单位总平面图", parent_key="images", blocks=[_image_block("单位总平面图", image_categories["单位总平面图"])]),
        DocumentPage(key="images_water", title="水源图", parent_key="images", blocks=[_image_block("水源图", image_categories["水源图"])]),
        DocumentPage(key="images_floor", title="楼层平面图", parent_key="images", blocks=[_image_block("楼层平面图", image_categories["楼层平面图"])]),
        DocumentPage(
            key="assist",
            title="辅助决策",
            blocks=[
                _rows_to_block("风险评估", risk_rows),
                _rows_to_block("灾情假设", scenario_rows),
                _rows_to_block("现场通信", communication_rows),
                _rows_to_block("首战力量部署", initial_rows),
                _image_block("首战力量部署图", image_categories["首战力量部署图"]),
                _rows_to_block("增援力量部署", reinforce_rows),
                _image_block("增援力量部署图", image_categories["增援力量部署图"]),
            ],
        ),
        DocumentPage(key="assist_scenario", title="灾情假设", parent_key="assist", blocks=[_rows_to_block("灾情假设", scenario_rows)]),
        DocumentPage(key="assist_communication", title="现场通信", parent_key="assist", blocks=[_rows_to_block("现场通信", communication_rows)]),
        DocumentPage(
            key="assist_initial",
            title="首战力量部署图",
            parent_key="assist",
            blocks=[_rows_to_block("首战力量部署", initial_rows), _image_block("首战力量部署图", image_categories["首战力量部署图"])],
        ),
        DocumentPage(
            key="assist_reinforce",
            title="增援力量部署图",
            parent_key="assist",
            blocks=[_rows_to_block("增援力量部署", reinforce_rows), _image_block("增援力量部署图", image_categories["增援力量部署图"])],
        ),
    ]

    navigation = [
        DocumentNavItem(
            key=item["key"],
            title=item["title"],
            children=[DocumentNavItem(**child) for child in item.get("children", [])],
        )
        for item in NAVIGATION
    ]

    return PlanDocumentResponse(
        plan_id=plan.id,
        plan_name=plan.name,
        pages=pages,
        navigation=navigation,
    )


def persist_plan_document(session: Session, plan: Plan, document_data: PlanDocumentResponse) -> None:
    session.query(PlanPage).filter(PlanPage.plan_id == plan.id).delete()
    session.flush()

    for page_index, page in enumerate(document_data.pages):
        page_model = PlanPage(
            plan_id=plan.id,
            key=page.key,
            title=page.title,
            parent_key=page.parent_key,
            sort_order=page_index,
        )
        session.add(page_model)
        session.flush()

        for block_index, block in enumerate(page.blocks):
            block_model = PlanBlock(
                page_id=page_model.id,
                type=block.type,
                title=block.title,
                sort_order=block_index,
                content=block.content,
            )
            session.add(block_model)
            session.flush()

            for row_index, row in enumerate(block.rows or []):
                for cell_order, cell in enumerate(row):
                    session.add(
                        PlanBlockCell(
                            block_id=block_model.id,
                            row_index=row_index,
                            cell_order=cell_order,
                            text=cell.text,
                            colspan=cell.colspan,
                            rowspan=cell.rowspan,
                            table_index=cell.table_index,
                            source_row_index=cell.row_index,
                            source_cell_index=cell.cell_index,
                            is_editable=cell.is_editable,
                        )
                    )

            for image_index, image in enumerate(block.images or []):
                session.add(
                    PlanBlockImage(
                        block_id=block_model.id,
                        image_name=image.image_name,
                        image_path=image.image_path,
                        doc_media_path=image.doc_media_path,
                        sort_order=image_index,
                    )
                )


def build_plan_document_from_db(session: Session, plan_id: int) -> PlanDocumentResponse:
    query = (
        session.query(Plan)
        .options(
            selectinload(Plan.pages)
            .selectinload(PlanPage.blocks)
            .selectinload(PlanBlock.cells),
            selectinload(Plan.pages)
            .selectinload(PlanPage.blocks)
            .selectinload(PlanBlock.images),
        )
        .filter(Plan.id == plan_id)
    )
    plan = query.one()

    if not plan.pages:
        images = session.query(PlanImage).filter(PlanImage.plan_id == plan.id).order_by(PlanImage.id.asc()).all()
        if not images and not plan.source_docx_path:
            raise FileNotFoundError("该预案尚未生成数据库主存内容")
        document_data = build_plan_document_from_source(plan, images)
        persist_plan_document(session, plan, document_data)
        session.commit()
        plan = query.one()

    pages: list[DocumentPage] = []
    for page in sorted(plan.pages, key=lambda item: item.sort_order):
        blocks: list[DocumentBlock] = []
        for block in sorted(page.blocks, key=lambda item: item.sort_order):
            rows: list[list[DocumentTableCell]] = []
            if block.type == "table":
                row_map: dict[int, list[DocumentTableCell]] = {}
                for cell in sorted(block.cells, key=lambda item: (item.row_index, item.cell_order)):
                    row_map.setdefault(cell.row_index, []).append(
                        DocumentTableCell(
                            text=cell.text,
                            colspan=cell.colspan,
                            rowspan=cell.rowspan,
                            block_id=block.id,
                            cell_order=cell.cell_order,
                            table_index=cell.table_index,
                            row_index=cell.row_index,
                            cell_index=cell.source_cell_index,
                            is_editable=cell.is_editable,
                        )
                    )
                rows = [row_map[index] for index in sorted(row_map)]
                for row in rows:
                    _apply_value_cell_editability(row)

            blocks.append(
                DocumentBlock(
                    type=block.type,
                    title=block.title,
                    rows=rows or None,
                    content=block.content,
                    images=list(block.images) or None,
                )
            )
        pages.append(
            DocumentPage(
                key=page.key,
                title=page.title,
                parent_key=page.parent_key,
                blocks=blocks,
            )
        )

    navigation = [
        DocumentNavItem(
            key=item["key"],
            title=item["title"],
            children=[DocumentNavItem(**child) for child in item.get("children", [])],
        )
        for item in NAVIGATION
    ]

    return PlanDocumentResponse(
        plan_id=plan.id,
        plan_name=plan.name,
        pages=pages,
        navigation=navigation,
    )
